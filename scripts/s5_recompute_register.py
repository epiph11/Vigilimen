#!/usr/bin/env python3
"""
Recompute the headline numbers of the 62443-3-2 risk register from its own inputs.

The workbook (docs/assessments/hydro-station/05-detailed-risk-assessment.xlsx) stores
formulas with NO cached values, so any viewer that does not recalculate (a previewer,
openpyxl data_only, a repository web view) shows empty bands. This script reads the
inputs and the lookup matrix, applies the same formulas, and prints what Excel would.

Not a CI check (needs openpyxl). It is how Station 5's claims were verified.

    pip install openpyxl
    python scripts/s5_recompute_register.py
"""
import collections
import sys
from pathlib import Path

import openpyxl

XLSX = Path(__file__).resolve().parent.parent / "docs/assessments/hydro-station/05-detailed-risk-assessment.xlsx"
wb = openpyxl.load_workbook(XLSX)
ws, sc = wb["Risk register"], wb["Scales"]

# Scales!B24:F28 = INDEX table, row = 6 - likelihood, column = consequence
matrix = {(L, C): sc.cell(24 + (5 - L), 1 + C).value for L in range(1, 6) for C in range(1, 6)}

rows = []
for r in range(5, 45):
    g = lambda col: ws[f"{col}{r}"].value            # noqa: E731
    C, Lu, Lr, cat = int(g("J")), int(g("K")), int(g("N")), g("I")
    res = matrix[(Lr, C)]
    if C == 5 and cat in ("Safety", "Environment"):
        treat = "life-safety override"
    elif res in ("E", "H"):
        treat = "YES"
    else:
        treat = "justify" if res == "M" else "no"
    rows.append(dict(id=g("A"), zone=g("B"), cat=cat, C=C, unm=matrix[(Lu, C)], res=res, treat=treat, prio=C * Lr))

ext = [x for x in rows if x["res"] == "E"]
print(f"risks: {len(rows)}")
print("residual bands:", dict(collections.Counter(x["res"] for x in rows)))
print("treatment flag:", dict(collections.Counter(x["treat"] for x in rows)))
print(f"flagged for treatment: {sum(x['treat'] in ('YES', 'life-safety override') for x in rows)} of {len(rows)}")
print(f"Extreme residual: {len(ext)}", [(x['id'], x['zone']) for x in ext])
print("Extreme in Z-01..Z-04:", [x["id"] for x in ext if x["zone"] in ("Z-01", "Z-02", "Z-03", "Z-04")] or "none")
print("top priority:", [(x["id"], x["prio"]) for x in sorted(rows, key=lambda x: -x["prio"])[:5]])
scored = sorted({x["zone"] for x in rows if x["zone"].startswith("C-")})
print("conduits with at least one risk row:", scored)
sys.exit(0)
